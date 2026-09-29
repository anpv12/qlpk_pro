from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from flask import Flask
from openpyxl import load_workbook

from app.api import medicine
from app.models.medicine import Medicine
from app.models.prescription import PrescriptionItem
from module_parts import setattr_all


def batch(identifier, imported, price, **extra):
    return SimpleNamespace(id=identifier, import_date=imported,
        created_at=extra.get('created_at', datetime.now(timezone.utc)),
        expiry_date=extra.get('expiry_date', imported + timedelta(days=365)),
        batch_number=f'LOT-{identifier}', import_price=price)


def test_latest_batch_uses_import_date_created_at_and_id():
    today = date.today()
    now = datetime.now(timezone.utc)
    older = batch(1, today - timedelta(days=1), 100, expiry_date=today + timedelta(days=900))
    no_timestamp = batch(4, today, 150, created_at=None)
    earlier = batch(5, today, 200, created_at=now - timedelta(seconds=1))
    same_time = batch(2, today, 300, created_at=now)
    latest = batch(3, today, 0, created_at=now)
    assert medicine.latest_batch_of([older, no_timestamp, earlier, latest, same_time]) is latest
    assert medicine.latest_batch_of([]) is None


@pytest.mark.parametrize('price', [None, Decimal('0'), Decimal('125.50')])
@pytest.mark.parametrize('excel', [False, True])
def test_inventory_report_preserves_fractional_stock_and_latest_price(monkeypatch, price, excel):
    today = date.today()
    older = batch(1, today - timedelta(days=1), 999, expiry_date=today + timedelta(days=900))
    latest = batch(2, today, price)
    catalog = SimpleNamespace(id=1, name='Test', internal_code='TEST', unit='viên',
        prescription_type='BASIC', stock_quantity=Decimal('10.5'), low_stock_threshold=0,
        unit_price=1000, import_price=888, batches=[older, latest])
    database = MagicMock()

    def query(*entities):
        result = MagicMock()
        for method in ['join', 'filter', 'group_by', 'options', 'order_by']:
            getattr(result, method).return_value = result
        if entities[0] is Medicine:
            result.all.return_value = [catalog]
        elif entities[0] is PrescriptionItem.medicine_id:
            result.all.return_value = [(1, Decimal('2.25'))]
        else:
            result.all.return_value = []
        return result

    database.query.side_effect = query
    setattr_all(monkeypatch,medicine, 'get_db', lambda: iter([database]))
    application = Flask(__name__)
    with application.test_request_context('/'):
        if excel:
            response = medicine.export_statistics_excel.__wrapped__(SimpleNamespace(id=1))
            assert response.status_code == 200
            workbook = load_workbook(BytesIO(response.get_data()))
            sheet = workbook['Tồn kho']
            assert sheet.cell(4, 4).value == (float(price) if price is not None else None)
            assert sheet.cell(4, 6).value == 2.25
            assert sheet.cell(4, 7).value == 10.5
            assert sheet.cell(4, 8).value == latest.expiry_date.strftime('%d/%m/%Y')
        else:
            response, status = medicine.get_statistics_inventory.__wrapped__(SimpleNamespace(id=1))
            assert status == 200, response.get_json()
            payload = response.get_json()
            row = payload['medicines'][0]
            assert row['stock_quantity'] == 10.5
            assert row['export_quantity'] == 2.25
            assert row['import_price'] == (float(price) if price is not None else None)
            assert row['batch_number'] == latest.batch_number
            assert payload['summary']['total_stock'] == 10.5
            assert payload['summary']['total_export'] == 2.25
