"""Read-only PostgreSQL checks for the bounded DAV autocomplete query."""
import os
import pytest
from flask import Flask
from sqlalchemy import event, text
from sqlalchemy.orm import Session
from app.core.database import engine
from app.modules.medicines.services.reference_catalog_query import list_reference_catalog

pytestmark = pytest.mark.skipif(os.environ.get('QLPK_RUN_DB_TESTS') != '1', reason='Opt-in read-only DB tests')


@pytest.fixture
def readonly():
    statements = []
    with engine.connect() as connection:
        transaction = connection.begin()
        connection.execute(text('SET TRANSACTION READ ONLY'))
        connection.execute(text("SET LOCAL statement_timeout = '15s'"))
        def capture(conn, cursor, statement, parameters, context, executemany):
            statements.append(statement)
        event.listen(connection, 'before_cursor_execute', capture)
        db = Session(bind=connection)
        try:
            yield db, statements
        finally:
            db.close()
            event.remove(connection, 'before_cursor_execute', capture)
            if transaction.is_active:
                transaction.rollback()


def test_lookup_is_bounded_and_matches_catalog_identity(readonly):
    db, statements = readonly
    result = list_reference_catalog(db, per_page=12, autocomplete=True)
    assert 0 < len(result['items']) <= 12
    assert result['total'] is None and result['total_pages'] is None
    assert len(statements) == 2
    assert not any('count(' in sql.lower() for sql in statements)
    projection = statements[0].lower().split('from')[0]
    assert 'raw_payload' not in projection
    normal = list_reference_catalog(db, per_page=12)
    assert result['has_more'] == (normal['total'] > 12)
    for lightweight, full in zip(result['items'], normal['items']):
        assert lightweight == {key: full[key] for key in lightweight}
        assert 'updated_at' in lightweight and 'clinic_medicine_id' in lightweight


def test_route_suggestion_matches_detail_and_clinic_without_writing_source(readonly):
    from app.modules.medicines.services.reference_catalog_query import get_reference_catalog_detail
    from app.models.medicine_reference_catalog import MedicineReferenceCatalog
    db, _ = readonly
    result = list_reference_catalog(db, registration_number='QLĐB1-H07-19', autocomplete=True)
    assert len(result['items']) == 1
    item = result['items'][0]
    detail = get_reference_catalog_detail(db, item['id'])
    assert item['route'] is None and detail['route'] is None
    assert item['suggested_route'] == detail['suggested_route'] == 'Tiêm'
    assert item['clinic_defaults']['suggested_administration_method'] == 'Tiêm'
    source = db.get(MedicineReferenceCatalog, item['id'])
    assert source.raw_payload['thongTinThuocCoBan'].get('tenDuongDung') is None
    assert not db.is_modified(source)


def test_lookup_pages_and_empty_search_results(readonly):
    db, _ = readonly
    first = list_reference_catalog(db, per_page=2, autocomplete=True)
    second = list_reference_catalog(db, page=2, per_page=2, autocomplete=True)
    assert first['has_more'] and second['items']
    assert not ({row['id'] for row in first['items']} & {row['id'] for row in second['items']})
    empty = list_reference_catalog(db, search='__QLPK_NO_DAV_MATCH_938540127__', autocomplete=True)
    assert empty['items'] == [] and empty['has_more'] is False


def test_autocomplete_api_skips_summary_but_default_contract_is_preserved(readonly, monkeypatch):
    import app.modules.medicines.api.reference_catalog as api
    db, statements = readonly
    monkeypatch.setattr(api, 'SessionLocal', lambda: db)
    app = Flask(__name__)
    with app.test_request_context('/api/medicine-reference-catalog?mode=autocomplete&per_page=12'):
        response = api.get_reference_catalog.__wrapped__(None)
        data = response.get_json()
    assert data['success'] and len(data['data']) == 12
    assert 'summary' not in data and 'total' not in data and 'has_more' in data
    assert len(statements) == 2
    statements.clear()
    with app.test_request_context('/api/medicine-reference-catalog?per_page=12'):
        response = api.get_reference_catalog.__wrapped__(None)
        data = response.get_json()
    assert data['success'] and 'summary' in data and 'total_pages' in data
    assert data['total'] >= len(data['data'])
    statements.clear()
    with app.test_request_context('/api/medicine-reference-catalog?search=tablet&include_summary=0&per_page=10'):
        response = api.get_reference_catalog.__wrapped__(None)
        data = response.get_json()
    assert data['success'] and data['summary'] is None and data['total'] > 0
    assert len(statements) == 3  # count + page + clinic links; no dashboard aggregates


def test_autocomplete_still_requires_authentication():
    from app.modules.medicines.api.reference_catalog import reference_catalog_bp
    app = Flask(__name__)
    app.register_blueprint(reference_catalog_bp)
    assert app.test_client().get('/api/medicine-reference-catalog?mode=autocomplete').status_code == 401
    assert app.test_client().get('/api/medicine-reference-catalog/export/excel').status_code == 401


@pytest.mark.parametrize('search,status', [('tablet', 'active'), ('VN-5518-10', 'all'), ('__NO_DAV_EXPORT__', 'active'), ('tablet', 'withdrawn')])
def test_export_matches_all_filtered_rows_and_preserves_workbook_types(readonly, search, status):
    from openpyxl import load_workbook
    from app.models.medicine_reference_catalog import MedicineReferenceCatalog as Catalog
    from app.modules.medicines.services.reference_catalog_query import build_reference_catalog_query
    from app.modules.medicines.services.reference_catalog_export import build_reference_catalog_excel
    db, _ = readonly
    expected = build_reference_catalog_query(db, search=search, status=status).order_by(
        Catalog.name, Catalog.registration_number, Catalog.id).all()
    sheet = load_workbook(build_reference_catalog_excel(db, search=search, status=status)).active
    assert sheet.auto_filter.ref == f'A6:O{6+len(expected)}'
    assert sheet.freeze_panes == 'D7'
    assert sheet['A6'].fill.fgColor.rgb == '00EFE2D5'
    assert sheet.max_row == len(expected) + 8
    for row, item in enumerate(expected, 7):
        assert sheet.cell(row, 2).value == item.source_id
        assert sheet.cell(row, 3).value == item.name
        assert sheet.cell(row, 4).value == item.active_ingredient
        assert sheet.cell(row, 11).value == item.registration_number
        if item.source_id:
            assert sheet.cell(row, 2).data_type == 's'
        if item.registration_expiry_date:
            assert sheet.cell(row, 14).value.date() == item.registration_expiry_date.date()
            assert sheet.cell(row, 14).number_format == 'dd/mm/yyyy'
