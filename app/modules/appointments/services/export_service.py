"""Excel export services for appointment workflows."""

from dataclasses import dataclass
from datetime import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.appointment import Appointment
from app.models.user import User
from app.utils.address_contract import build_full_address
from app.utils.allergy_contract import format_allergy_entries
from app.utils.examination_utils import current_examination
from app.utils.risk_assessment import format_risk_assessment


@dataclass
class AppointmentExportResult:
    output: io.BytesIO
    filename: str
    mimetype: str


def build_appointments_export_file(db, data, logger=None):
    """Build the legacy appointment-management Excel export file."""
    appointments = _get_export_appointments(db, data)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Lịch hẹn'

    styles = _build_styles()
    headers = [
        'STT', 'Ngày hẹn', 'Giờ hẹn', 'Bác sĩ / TLG', 'Trạng thái', 'Loại hẹn',
        'Họ và tên', 'Số điện thoại', 'CCCD', 'Email', 'Ngày sinh', 'Giới tính',
        'Địa chỉ', 'Dịch vụ / Gói khám', 'Loại khám', 'Tiền căn bệnh', 'Dị ứng',
        'Thuốc đang dùng', 'Lý do chính', 'Triệu chứng', 'Ghi chú'
    ]
    col_widths = [5, 12, 8, 30, 16, 12, 25, 15, 15, 25, 12, 8, 35, 30, 14, 25, 20, 20, 25, 25, 30]

    _write_header(ws, headers, col_widths, styles)
    _write_appointment_rows(db, ws, appointments, styles, logger=logger)

    ws.freeze_panes = 'A4'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return AppointmentExportResult(
        output=output,
        filename=f'lich_hen_{datetime.now().strftime("%Y%m%d")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def _get_export_appointments(db, data):
    doctor_id = data.get('doctor_id')
    role_filter = data.get('role_filter')

    query = db.query(Appointment).filter(Appointment.is_deleted == False)

    if doctor_id:
        query = query.filter(
            (Appointment.doctor_id == int(doctor_id)) |
            (Appointment.psychologist_id == int(doctor_id))
        )

    if role_filter:
        role_users = db.query(User.id).filter(User.role == role_filter.upper()).all()
        role_ids = [u.id for u in role_users]
        if role_ids:
            query = query.filter(
                (Appointment.doctor_id.in_(role_ids)) |
                (Appointment.psychologist_id.in_(role_ids))
            )

    return query.order_by(Appointment.appointment_date.desc()).all()


def _build_styles():
    thin_border = Border(
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB')
    )
    header_border = Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000'),
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000')
    )

    return {
        'title_font': Font(name='Arial', bold=True, size=14, color='1A3C5E'),
        'header_font': Font(name='Arial', bold=True, size=10, color='FFFFFF'),
        'header_fill': PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid'),
        'header_align': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'data_font': Font(name='Arial', size=10),
        'data_align': Alignment(vertical='center', wrap_text=True),
        'center_align': Alignment(horizontal='center', vertical='center'),
        'thin_border': thin_border,
        'header_border': header_border,
    }


def _write_header(ws, headers, col_widths, styles):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=f'DANH SÁCH LỊCH HẸN — Xuất ngày {datetime.now().strftime("%d/%m/%Y")}')
    title_cell.font = styles['title_font']
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 6

    for ci, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=ci, value=header)
        cell.font = styles['header_font']
        cell.fill = styles['header_fill']
        cell.alignment = styles['header_align']
        cell.border = styles['header_border']
    ws.row_dimensions[3].height = 24

    for ci, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = width


def _write_appointment_rows(db, ws, appointments, styles, logger=None):
    for idx, appt in enumerate(appointments, 1):
        row_num = idx + 3
        row_data = _build_appointment_export_row(db, idx, appt, logger=logger)

        for ci, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=ci, value=value)
            cell.font = styles['data_font']
            cell.alignment = styles['center_align'] if ci == 1 else styles['data_align']
            cell.border = styles['thin_border']


def _build_appointment_export_row(db, idx, appt, logger=None):
    patient = appt.patient
    doctor = appt.doctor

    status_map = {'SCHEDULED': 'Chờ xác nhận', 'CONFIRMED': 'Đã xác nhận', 'NO_SHOW': 'Không đến', 'CANCELLED': 'Hủy', 'COMPLETED': 'Hoàn thành'}
    cat_map = {'RE_EXAMINATION': 'Tái khám', 'NEW': 'Khám mới'}
    type_map = {'SERVICE': 'Theo dịch vụ', 'PACKAGE': 'Theo gói'}

    ngay_hen = appt.appointment_date.strftime('%d/%m/%Y') if appt.appointment_date else ''
    gio_hen = appt.appointment_date.strftime('%H:%M') if appt.appointment_date else ''
    bac_si = doctor.full_name if doctor else ''

    status_val = appt.status.value if appt.status else ''
    trang_thai = status_map.get(status_val, status_val)

    cat_val = appt.appointment_category.value if appt.appointment_category else ''
    loai_hen = cat_map.get(cat_val, cat_val)

    type_val = appt.appointment_type.value if appt.appointment_type else ''
    loai_kham = type_map.get(type_val, type_val)

    ho_ten = patient.full_name if patient else ''
    sdt = patient.phone if patient else ''
    cccd = patient.id_number if patient and hasattr(patient, 'id_number') else ''
    email = patient.email if patient and hasattr(patient, 'email') else ''

    ngay_sinh = ''
    if patient and patient.date_of_birth:
        ngay_sinh = patient.date_of_birth.strftime('%d/%m/%Y') if hasattr(patient.date_of_birth, 'strftime') else str(patient.date_of_birth)

    gioi_tinh = ''
    if patient and hasattr(patient, 'gender') and patient.gender:
        gender_value = patient.gender.value if hasattr(patient.gender, 'value') else str(patient.gender)
        gioi_tinh = {'MALE': 'Nam', 'FEMALE': 'Nữ', 'OTHER': 'Khác'}.get(gender_value, gender_value)

    dia_chi = ''
    if patient:
        dia_chi = build_full_address(
            getattr(patient, 'address', None) or getattr(patient, 'address_detail', None),
            getattr(patient, 'ward', None),
            getattr(patient, 'district', None),
            getattr(patient, 'province', None)
        )

    dich_vu = ''
    if appt.service:
        dich_vu = appt.service.name
    elif hasattr(appt, 'package') and appt.package:
        dich_vu = appt.package.name

    tien_can = _build_history_text(db, patient, appt, logger=logger)
    di_ung = format_allergy_entries(getattr(patient, 'allergies', []) if patient else [])
    thuoc = (getattr(patient, 'current_medication', '') or '') if patient else ''

    ly_do = ''
    trieu_chung = ''
    exam = current_examination(appt)
    if exam:
        ly_do = getattr(exam, 'main_reason', '') or ''
        trieu_chung = getattr(exam, 'main_symptoms', '') or ''

    return [
        idx, ngay_hen, gio_hen, bac_si, trang_thai, loai_hen,
        ho_ten, sdt, cccd or '', email or '', ngay_sinh, gioi_tinh,
        dia_chi, dich_vu, loai_kham, tien_can, di_ung,
        thuoc, ly_do, trieu_chung, appt.notes or ''
    ]


def _build_history_text(db, patient, appt, logger=None):
    if not patient:
        return ''

    tien_can_parts = []
    if patient.physical_history:
        try:
            from app.models.icd import ICD
            icd_ids = []
            text_vals = []
            if isinstance(patient.physical_history, list):
                for item in patient.physical_history:
                    if isinstance(item, dict):
                        if item.get('type') == 'icd' and item.get('id') is not None:
                            icd_ids.append(item.get('id'))
                        elif item.get('type') == 'text' and item.get('value'):
                            text_vals.append(item.get('value'))
                    elif isinstance(item, int):
                        icd_ids.append(item)

            if icd_ids:
                icd_list = db.query(ICD).filter(ICD.id.in_(icd_ids)).all()
                for item in icd_list:
                    tien_can_parts.append(f"{item.icd_code} - {item.disease_name}")
            if text_vals:
                tien_can_parts.extend(text_vals)
        except Exception as exc:
            if logger:
                logger.error(f"Error fetching ICD names in excel export: {exc}", exc_info=True)

    exam_risk = current_examination(appt)
    if exam_risk and getattr(exam_risk, 'risk_assessment', None):
        tien_can_parts.append(format_risk_assessment(exam_risk.risk_assessment))

    return "; ".join(tien_can_parts)
