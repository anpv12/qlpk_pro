"""Read-only DAV export using the application's existing openpyxl runtime."""
from datetime import datetime
from io import BytesIO
from math import ceil
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.medicine_reference_catalog import MedicineReferenceCatalog as Catalog
from app.modules.medicines.services.reference_catalog_query import build_reference_catalog_query


STATUS_LABELS = {
    'active': 'Đang hiệu lực', 'expired': 'Hết hạn', 'withdrawn': 'Đã rút số',
    'deleted': 'Đã xóa', 'all': 'Tất cả',
}
COLUMNS = (
    ('STT', None, 7), ('Mã nguồn DAV', 'source_id', 17), ('Tên thuốc', 'name', 42),
    ('Hoạt chất', 'active_ingredient', 42), ('Hàm lượng', 'strength', 28),
    ('Dạng bào chế', 'dosage_form', 25), ('Quy cách đóng gói', 'packaging', 42),
    ('Đường dùng', 'route', 22), ('Nhà sản xuất', 'manufacturer_name', 42),
    ('Nước sản xuất', 'manufacturer_country', 22),
    ('Số đăng ký', 'registration_number', 23),
    ('Số đăng ký cũ', 'old_registration_number', 23),
    ('Ngày cấp SĐK', 'registration_issue_date', 19),
    ('Ngày hết hạn SĐK', 'registration_expiry_date', 21),
    ('Trạng thái', 'status', 20),
)


def _status(item):
    if item.withdrawn:
        return 'Đã rút số'
    if item.is_deleted:
        return 'Đã xóa'
    if item.is_expired:
        return 'Hết hạn'
    return 'Đang hiệu lực' if item.is_active else 'Ngừng hiệu lực'


def build_reference_catalog_excel(db, search='', status='active'):
    """Export every matching row in stable list order, with bounded row memory."""
    status = status if status in STATUS_LABELS else 'active'
    search = (search or '').strip()
    query = build_reference_catalog_query(db, search=search, status=status)
    fields = [getattr(Catalog, field) for _, field, _ in COLUMNS if field not in (None, 'status')]
    query = query.with_entities(
        *fields, Catalog.is_active, Catalog.is_expired, Catalog.is_deleted,
        Catalog.raw_payload['isDaRutSoDangKy'].as_boolean().label('withdrawn'),
    ).order_by(Catalog.name, Catalog.registration_number, Catalog.id)

    workbook = Workbook(write_only=True)
    workbook.properties.creator = 'Sơn Tâm Clinic'
    workbook.properties.title = 'Danh mục thuốc DAV'
    sheet = workbook.create_sheet('Danh mục DAV')
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'D7'
    sheet.print_title_rows = '1:6'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = 'Sơn Tâm Clinic — Trang &P / &N'
    for index, (_, _, width) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    normal_font = Font(name='Calibri', size=11, color='342D28')
    label_font = Font(name='Calibri', size=11, bold=True, color='5B3829')
    title_font = Font(name='Calibri', size=18, bold=True, color='5B3829')
    cream = PatternFill('solid', fgColor='F8F3EC')
    header_fill = PatternFill('solid', fgColor='EFE2D5')
    alignment = Alignment(vertical='top', wrap_text=True)

    def cell(value, font=normal_font, fill=None):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub('', value)
        result = WriteOnlyCell(sheet, value=value)
        # External names/codes are literal text, including leading =, + or -.
        if isinstance(value, str):
            result.data_type = 's'
            result.number_format = '@'
        if isinstance(value, datetime):
            result.value = value.date()
            result.number_format = 'dd/mm/yyyy'
        result.font = font
        result.alignment = alignment
        if fill:
            result.fill = fill
        return result

    stamp = datetime.now(ZoneInfo('Asia/Ho_Chi_Minh')).strftime('%d/%m/%Y %H:%M')
    titles = (
        ('SƠN TÂM CLINIC', label_font, 24),
        ('DANH MỤC THUỐC DAV', title_font, 32),
        ('Nguồn: danh mục DAV đã đồng bộ trong ứng dụng. Không phải tồn kho phòng khám.', normal_font, 24),
        (f'Xuất lúc: {stamp}   |   Hiệu lực: {STATUS_LABELS[status]}   |   Từ khóa: {search or "Tất cả"}', normal_font, 32),
    )
    last_column = get_column_letter(len(COLUMNS))
    for index, (title, font, height) in enumerate(titles, 1):
        sheet.row_dimensions[index].height = height
        sheet.merged_cells.add(f'A{index}:{last_column}{index}')
        sheet.append([cell(title, font, cream)] + [cell(None, fill=cream) for _ in COLUMNS[1:]])
    sheet.append([])
    sheet.row_dimensions[6].height = 32
    sheet.append([cell(label, label_font, header_fill) for label, _, _ in COLUMNS])

    count = 0
    for count, item in enumerate(query.yield_per(500), 1):
        values = [count if field is None else _status(item) if field == 'status' else getattr(item, field)
                  for _, field, _ in COLUMNS]
        lines = max(sum(max(1, ceil(len(line) / (width - 3))) for line in str(value or '').split('\n'))
                    if isinstance(value, str) else 1 for value, (_, _, width) in zip(values, COLUMNS))
        sheet.row_dimensions[6 + count].height = min(409, max(30, lines * 15 + 8))
        sheet.append([cell(value, fill=cream if count % 2 == 0 else None) for value in values])
    sheet.auto_filter.ref = f'A6:{last_column}{6 + count}'
    sheet.append([])
    sheet.append([cell(f'Tổng số bản ghi: {count:,}'.replace(',', '.'), label_font)])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
