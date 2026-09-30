"""app.api.dashboard: phần 4 — tách từ dashboard.py (import ở cuối dashboard.py để đăng ký route/giữ tên cũ)."""

from flask import jsonify, request
from sqlalchemy import func
from app.core.database import get_db
from app.models.appointment import Appointment
from app.models.user import User, UserRole
from app.api.auth import require_auth
from app.api.dashboard import (  # noqa: E402 — module gốc đã khởi tạo xong các tên này
    dashboard_bp,
    logger,
)


def _write_thu_chi_total_row(brd, daily_data, right_align, total_chi, total_diff, total_fill, total_med, total_svc, total_thu, ws):
    from openpyxl.styles import Font, Alignment
    # Total row
    tr = len(daily_data) + 4
    ws.cell(row=tr, column=1, value="")
    ws.cell(row=tr, column=2, value="TỔNG CỘNG").font = Font(name='Arial', bold=True, size=11, color="0F766E")
    for c, v in enumerate([total_svc, total_med, total_thu, total_chi, total_diff], 3):
        cell = ws.cell(row=tr, column=c, value=v)
        cell.font = Font(name='Arial', bold=True, size=11, color="0F766E")
        cell.number_format = '#,##0'
        cell.alignment = right_align
        cell.border = brd
        cell.fill = total_fill
    ws.cell(row=tr, column=2).fill = total_fill
    ws.cell(row=tr, column=2).border = brd
    ws.cell(row=tr, column=2).alignment = Alignment(vertical='center')
    ws.cell(row=tr, column=1).fill = total_fill
    ws.cell(row=tr, column=1).border = brd

    # Diff total color
    diff_total_cell = ws.cell(row=tr, column=7)
    diff_total_cell.font = Font(name='Arial', bold=True, size=11, color="059669" if total_diff >= 0 else "DC2626")


def _write_thu_chi_daily_rows(brd, center_align, daily_data, data_font, right_align, ws):
    from openpyxl.styles import Font, Alignment
    # Data rows
    total_svc = total_med = total_thu = total_chi = total_diff = 0
    for i, row in enumerate(daily_data, 1):
        r = i + 3
        data = [
            i,
            row['date'].strftime('%d/%m/%Y'),
            row['service'],
            row['medicine'],
            row['thu'],
            row['chi'],
            row['diff']
        ]
        for c, v in enumerate(data, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = brd
            cell.font = data_font
            cell.alignment = Alignment(vertical='center')
            if c >= 3:
                cell.number_format = '#,##0'
                cell.alignment = right_align
            if c == 1:
                cell.alignment = center_align
            if c == 2:
                cell.alignment = center_align

        # Color the diff column
        diff_cell = ws.cell(row=r, column=7)
        if row['diff'] >= 0:
            diff_cell.font = Font(name='Arial', color="059669", bold=True)
        else:
            diff_cell.font = Font(name='Arial', color="DC2626", bold=True)

        total_svc += row['service']
        total_med += row['medicine']
        total_thu += row['thu']
        total_chi += row['chi']
        total_diff += row['diff']
    return total_chi, total_diff, total_med, total_svc, total_thu


def _collect_thu_chi_daily(db, end_date, start_date):
    from datetime import timedelta, datetime as dt
    from app.models.prescription import Prescription, PrescriptionItem
    from app.models.expense import Expense
    # --- Query revenue per day ---
    def get_service_revenue_daily(d):
        from app.models.service import Service
        s = dt.combine(d, dt.min.time())
        e = dt.combine(d, dt.max.time())
        result = db.query(func.coalesce(func.sum(Service.default_price), 0))\
            .join(Appointment, Appointment.service_id == Service.id)\
            .filter(Appointment.appointment_date >= s, Appointment.appointment_date <= e,
                    Appointment.status == 'CONFIRMED', Appointment.is_deleted == False).scalar() or 0
        return float(result)

    def get_medicine_revenue_daily(d):
        s = dt.combine(d, dt.min.time())
        e = dt.combine(d, dt.max.time())
        result = db.query(func.coalesce(func.sum(PrescriptionItem.unit_price * PrescriptionItem.quantity), 0))\
            .join(Prescription, PrescriptionItem.prescription_id == Prescription.id)\
            .join(Appointment, Prescription.appointment_id == Appointment.id)\
            .filter(Appointment.appointment_date >= s, Appointment.appointment_date <= e).scalar() or 0
        return float(result)

    # --- Query expense per day ---
    expense_rows = db.query(
        Expense.date,
        func.sum(Expense.amount).label('total')
    ).filter(
        Expense.date >= start_date,
        Expense.date <= end_date
    ).group_by(Expense.date).all()

    chi_by_date = {r.date: float(r.total or 0) for r in expense_rows}

    # --- Build daily data ---
    daily_data = []
    d = start_date
    while d <= end_date:
        svc = get_service_revenue_daily(d)
        med = get_medicine_revenue_daily(d)
        thu = svc + med
        chi = chi_by_date.get(d, 0)
        if thu > 0 or chi > 0:
            daily_data.append({
                'date': d,
                'service': svc,
                'medicine': med,
                'thu': thu,
                'chi': chi,
                'diff': thu - chi
            })
        d += timedelta(days=1)
    return daily_data


def _write_thu_chi_header(brd, center_align, header_fill_blue, header_font, period_label, ws):
    from openpyxl.styles import Font, Alignment
    # Title
    ws.merge_cells('A1:G1')
    title_cell = ws['A1']
    title_cell.value = f"THỐNG KÊ THU CHI — {period_label}"
    title_cell.font = Font(name='Arial', bold=True, size=14, color="0F766E")
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Blank row
    ws.append([])

    # Headers row 3
    headers = ["STT", "Ngày", "Dịch vụ (đ)", "Thuốc (đ)", "Tổng thu (đ)", "Chi tiêu (đ)", "Chênh lệch (đ)"]
    for c, v in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=v)
        cell.font = header_font
        cell.fill = header_fill_blue
        cell.alignment = center_align
        cell.border = brd
    ws.row_dimensions[3].height = 24


@dashboard_bp.route('/api/dashboard/export-thu-chi', methods=['GET'])
@require_auth
def export_thu_chi_excel(user):
    """Xuất Excel thống kê thu chi theo ngày"""
    import io
    from datetime import date, timedelta, datetime as dt
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from flask import send_file
    except ImportError:
        return jsonify({'error': 'openpyxl chưa được cài đặt'}), 500

    db = next(get_db())
    try:
        today = date.today()
        from_str = request.args.get('from_date', '')
        to_str = request.args.get('to_date', '')

        if from_str and to_str:
            start_date = dt.strptime(from_str, '%Y-%m-%d').date()
            end_date = dt.strptime(to_str, '%Y-%m-%d').date()
        else:
            start_date = today - timedelta(days=30)
            end_date = today

        # Cap to today
        if end_date > today:
            end_date = today

        period_label = f"{start_date.strftime('%d/%m/%Y')} – {end_date.strftime('%d/%m/%Y')}"

        # --- Styles ---
        thin = Side(style='thin', color='000000')
        brd = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill_blue = PatternFill("solid", fgColor="0F766E")
        header_font = Font(name='Arial', bold=True, color="FFFFFF", size=11)
        data_font = Font(name='Arial', size=10)
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        right_align = Alignment(horizontal='right', vertical='center')
        total_fill = PatternFill("solid", fgColor="F0FDF4")

        daily_data = _collect_thu_chi_daily(db, end_date, start_date)

        # --- Build workbook ---
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Thu Chi"

        _write_thu_chi_header(brd, center_align, header_fill_blue, header_font, period_label, ws)

        total_chi, total_diff, total_med, total_svc, total_thu = _write_thu_chi_daily_rows(brd, center_align, daily_data, data_font, right_align, ws)

        _write_thu_chi_total_row(brd, daily_data, right_align, total_chi, total_diff, total_fill, total_med, total_svc, total_thu, ws)

        # Column widths
        for col, w in [('A', 6), ('B', 14), ('C', 16), ('D', 16), ('E', 18), ('F', 16), ('G', 18)]:
            ws.column_dimensions[col].width = w

        # Freeze header
        ws.freeze_panes = 'A4'

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"thong_ke_thu_chi_{from_str or today.strftime('%Y-%m-%d')}_{to_str or today.strftime('%Y-%m-%d')}.xlsx"
        return send_file(output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name=filename)

    except Exception as e:
        logger.error(f"Export thu chi excel error: {e}", exc_info=True)
        return jsonify({'error': str(e)}), 500
    finally:
        db.close()


@dashboard_bp.route('/api/dashboard/exam-stats-by-day', methods=['GET'])
@require_auth
def get_exam_stats_by_day(user):
    """Số lượt khám theo từng ngày trong khoảng from_date → to_date, phân nhóm theo role"""
    db = next(get_db())
    try:
        from datetime import date, timedelta, datetime as dt
        today = date.today()
        from_str = request.args.get('from_date', (today - timedelta(days=29)).isoformat())
        to_str   = request.args.get('to_date',   today.isoformat())
        start_date = dt.strptime(from_str, '%Y-%m-%d').date()
        end_date   = dt.strptime(to_str,   '%Y-%m-%d').date()

        start_dt = dt.combine(start_date, dt.min.time())
        end_dt   = dt.combine(end_date,   dt.max.time())

        # Query tất cả appointments trong khoảng, join User để lấy role
        rows = db.query(
            func.date(Appointment.appointment_date).label('day'),
            User.role.label('doctor_role'),
            func.count(Appointment.id).label('cnt')
        ).join(
            User, Appointment.doctor_id == User.id
        ).filter(
            Appointment.appointment_date >= start_dt,
            Appointment.appointment_date <= end_dt,
            Appointment.is_deleted == False,
            Appointment.status.notin_(['CANCELLED', 'NO_SHOW'])
        ).group_by(
            func.date(Appointment.appointment_date),
            User.role
        ).all()

        # Build day_map: { 'YYYY-MM-DD': { 'doctor': N, 'psychologist': N } }
        day_map = {}
        for r in rows:
            key = str(r.day)
            if key not in day_map:
                day_map[key] = {'doctor': 0, 'psychologist': 0}
            role_val = r.doctor_role.value.lower() if hasattr(r.doctor_role, 'value') else str(r.doctor_role).lower()
            if role_val == UserRole.PSYCHOLOGIST.value.lower():
                day_map[key]['psychologist'] += int(r.cnt)
            else:
                day_map[key]['doctor'] += int(r.cnt)

        items = []
        d = start_date
        while d <= end_date:
            key = d.isoformat()
            counts = day_map.get(key, {'doctor': 0, 'psychologist': 0})
            items.append({
                'label': f'{d.day}/{d.month}',
                'date': key,
                'count': counts['doctor'] + counts['psychologist'],
                'doctor_count': counts['doctor'],
                'psychologist_count': counts['psychologist']
            })
            d += timedelta(days=1)

        return jsonify({'items': items}), 200

    except Exception as e:
        logger.error(f"Error getting exam stats by day: {e}")
        return jsonify({'error': str(e)}), 500
    finally:
        db.close()


@dashboard_bp.route('/api/dashboard/exam-detail-by-day', methods=['GET'])
@require_auth
def get_exam_detail_by_day(user):
    """Chi tiết ca khám trong 1 ngày cụ thể"""
    db = next(get_db())
    try:
        from datetime import datetime as dt
        from app.models.patient import Patient
        from app.models.service import Service

        date_str = request.args.get('date', '')
        if not date_str:
            return jsonify({'error': 'Thiếu tham số date'}), 400

        target = dt.strptime(date_str, '%Y-%m-%d')
        start  = dt.combine(target.date(), dt.min.time())
        end    = dt.combine(target.date(), dt.max.time())

        rows = db.query(
            Appointment, Patient, User
        ).join(Patient, Appointment.patient_id == Patient.id)\
         .join(User,    Appointment.doctor_id  == User.id)\
         .filter(
            Appointment.appointment_date >= start,
            Appointment.appointment_date <= end,
            Appointment.is_deleted == False,
            Appointment.status.notin_(['CANCELLED', 'NO_SHOW'])
         ).order_by(Appointment.appointment_date).all()

        items = []
        for appt, patient, doctor in rows:
            svc_name = ''
            if appt.service_id:
                svc = db.query(Service).filter(Service.id == appt.service_id).first()
                if svc:
                    svc_name = svc.name

            status_map = {
                'SCHEDULED': 'Chờ xác nhận', 'CONFIRMED': 'Đã xác nhận',
                'COMPLETED': 'Hoàn thành',   'IN_PROGRESS': 'Đang khám',
            }
            items.append({
                'time':         appt.appointment_date.strftime('%H:%M'),
                'patient_name': patient.full_name,
                'phone':        getattr(patient, 'phone_number', '') or '',
                'doctor_name':  doctor.full_name,
                'doctor_role':  doctor.role.value if doctor.role else '',
                'service':      svc_name,
                'status':       status_map.get(str(appt.status).replace('AppointmentStatus.', ''), str(appt.status)),
                'category':     appt.appointment_category.value if appt.appointment_category else '',
                'notes':        appt.notes or '',
            })

        return jsonify({'date': date_str, 'items': items, 'total': len(items)}), 200

    except Exception as e:
        logger.error(f"Error getting exam detail by day: {e}")
        return jsonify({'error': str(e)}), 500
    finally:
        db.close()
