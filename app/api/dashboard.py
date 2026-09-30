"""Dashboard API — Revenue & ICD stats"""
from flask import Blueprint, jsonify, request
from sqlalchemy import String, cast, func, text
from app.core.database import get_db
from app.models.examination import Examination
from app.models.appointment import Appointment
from app.models.medicine import Medicine
from app.models.patient import Patient
from app.api.auth import require_auth
from app.utils.referral_source import (
    REFERRAL_SOURCE_BY_KEY,
    get_referral_source_key,
)
import logging
from app.utils.api_error_contract import api_error_boundary

logger = logging.getLogger(__name__)

dashboard_bp = Blueprint('dashboard_api', __name__)

def _get_patient_referral_source_key(patient):
    return get_referral_source_key(patient)


def _parse_dashboard_date_range(default_days=6):
    from datetime import date, timedelta, datetime as dt

    today = date.today()
    from_str = request.args.get('from_date', (today - timedelta(days=default_days)).isoformat())
    to_str = request.args.get('to_date', today.isoformat())
    start = dt.strptime(from_str, '%Y-%m-%d')
    end = dt.combine(dt.strptime(to_str, '%Y-%m-%d').date(), dt.max.time())
    return from_str, to_str, start, end

def _parse_referral_source_date_range():
    from datetime import date, datetime as dt

    today = date.today()
    year_start = date(today.year, 1, 1)
    from_str = request.args.get('from_date', year_start.isoformat())
    to_str = request.args.get('to_date', today.isoformat())
    start = dt.strptime(from_str, '%Y-%m-%d')
    end = dt.combine(dt.strptime(to_str, '%Y-%m-%d').date(), dt.max.time())
    return from_str, to_str, start, end


def _get_referral_source_cases(db, start, end):
    rows = db.query(Appointment, Patient, Examination)\
        .join(Patient, Appointment.patient_id == Patient.id)\
        .outerjoin(Examination, Examination.appointment_id == Appointment.id)\
        .filter(
            Appointment.appointment_date >= start,
            Appointment.appointment_date <= end,
            Appointment.is_deleted == False,
            Appointment.status.notin_(['CANCELLED', 'NO_SHOW'])
        )\
        .order_by(Appointment.appointment_date.desc())\
        .all()

    cases_by_appointment = {}
    for appt, patient, exam in rows:
        source_raw = (patient.referral_source_detail or patient.referral_source or '').strip()
        source_key = _get_patient_referral_source_key(patient)
        if not source_key:
            continue

        case = cases_by_appointment.get(appt.id)
        if not case:
            source_cfg = REFERRAL_SOURCE_BY_KEY[source_key]
            case = {
                'appointment_id': appt.id,
                'patient_id': patient.id,
                'patient_name': patient.full_name or '',
                'phone': patient.phone or '',
                'exam_date': appt.appointment_date.strftime('%d/%m/%Y') if appt.appointment_date else '',
                'exam_time': appt.appointment_date.strftime('%H:%M') if appt.appointment_date else '',
                'main_reason': '',
                'source_key': source_key,
                'source_label': source_cfg['label'],
                'source_value': source_raw or source_cfg['label'],
            }
            cases_by_appointment[appt.id] = case

        if exam and exam.main_reason and not case['main_reason']:
            case['main_reason'] = exam.main_reason

    return list(cases_by_appointment.values())

@dashboard_bp.route('/api/dashboard/service-summary', methods=['GET'])
@require_auth
@api_error_boundary(error='{error}')
def get_service_summary(user):
    """Tổng hợp dịch vụ / thuốc theo khoảng thời gian đang chọn"""
    db = next(get_db())
    try:
        from datetime import date, timedelta, datetime as dt
        from app.models.service import Service
        from app.models.prescription import Prescription, PrescriptionItem

        mode = request.args.get('mode', 'week')
        year_str = request.args.get('year', '')

        # Build time bounds — support from_date/to_date or legacy mode
        today = date.today()
        from_str = request.args.get('from_date', '')
        to_str = request.args.get('to_date', '')
        if from_str and to_str:
            start = dt.strptime(from_str, '%Y-%m-%d')
            end = dt.combine(dt.strptime(to_str, '%Y-%m-%d').date(), dt.max.time())
        elif mode == 'month':
            target_year = int(year_str) if year_str else today.year
            start = dt(target_year, 1, 1)
            end = dt(target_year, 12, 31, 23, 59, 59)
        else:  # week
            week_start = today - timedelta(days=6)
            start = dt.combine(week_start, dt.min.time())
            end = dt.combine(today, dt.max.time())

        # --- Service ranking ---
        svc_rows = db.query(
            Service.name.label('name'),
            func.count(Appointment.id).label('count'),
            func.sum(Service.default_price).label('revenue')
        ).join(
            Appointment, Appointment.service_id == Service.id
        ).filter(
            Appointment.appointment_date >= start,
            Appointment.appointment_date <= end,
            Appointment.status == 'CONFIRMED',
            Appointment.is_deleted == False
        ).group_by(Service.name).order_by(func.sum(Service.default_price).desc()).all()

        # --- Medicine ranking ---
        medicine_group_key = func.coalesce(cast(PrescriptionItem.medicine_id, String), PrescriptionItem.medicine_name)
        med_rows = db.query(
            func.coalesce(func.max(Medicine.name), func.max(PrescriptionItem.medicine_name)).label('name'),
            func.sum(PrescriptionItem.quantity).label('count'),
            func.sum(PrescriptionItem.unit_price * PrescriptionItem.quantity).label('revenue')
        ).join(
            Prescription, PrescriptionItem.prescription_id == Prescription.id
        ).join(
            Appointment, Prescription.appointment_id == Appointment.id
        ).outerjoin(
            Medicine, PrescriptionItem.medicine_id == Medicine.id
        ).filter(
            Appointment.appointment_date >= start,
            Appointment.appointment_date <= end
        ).group_by(medicine_group_key).order_by(
            func.sum(PrescriptionItem.unit_price * PrescriptionItem.quantity).desc()
        ).all()

        return jsonify({
            'mode': mode,
            'services': [{'name': r.name, 'count': int(r.count), 'revenue': float(r.revenue)} for r in svc_rows],
            'medicines': [{'name': r.name, 'count': float(r.count), 'revenue': float(r.revenue)} for r in med_rows],
        }), 200

    finally:
        db.close()

def _write_icd_case_groups(detail_hdr_fill, grouped, grp_fill, mk_border, row_num, sorted_codes, w_fill, ws2, z_fill):
    from openpyxl.styles import Font, Alignment
    for code in sorted_codes:
        grp = grouped[code]
        case_count = len(grp['cases'])

        # Group header row
        for c, v in enumerate([f"{code} — {grp['name']}", '', '', f"{case_count} ca", ''], 1):
            cell = ws2.cell(row=row_num, column=c, value=v)
            cell.font = Font(bold=True, size=11, color="5B21B6")
            cell.fill = grp_fill
            cell.border = mk_border()
        ws2.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
        row_num += 1

        # Detail header
        for c, v in enumerate(["STT", "Bệnh nhân", "Bác sĩ", "Ngày khám", "Giờ"], 1):
            cell = ws2.cell(row=row_num, column=c, value=v)
            cell.font = Font(bold=True, size=10, color="64748B")
            cell.fill = detail_hdr_fill
            cell.border = mk_border()
            cell.alignment = Alignment(horizontal='center', vertical='center')
        row_num += 1

        # Detail rows
        for j, case in enumerate(grp['cases'], 1):
            fill = z_fill if j % 2 == 0 else w_fill
            for c, v in enumerate([j, case['patient'], case['doctor'], case['date'], case['time']], 1):
                cell = ws2.cell(row=row_num, column=c, value=v)
                cell.border = mk_border()
                cell.fill = fill
                cell.alignment = Alignment(vertical='center')
            row_num += 1

        row_num += 1  # blank row between groups


def _group_icd_cases(db, end, start):
    from app.models.icd import ICD
    from app.models.patient import Patient
    from app.models.user import User
    all_exams = db.query(
        Examination, Patient, User, Appointment
    ).join(Appointment, Examination.appointment_id == Appointment.id)\
     .join(Patient, Examination.patient_id == Patient.id)\
     .join(User, Examination.doctor_id == User.id)\
     .filter(
        Examination.diagnosis.isnot(None), 
        Examination.diagnosis != text("'[]'::jsonb"),
        Appointment.appointment_date >= start, Appointment.appointment_date <= end
     ).order_by(Appointment.appointment_date.desc()).all()

    icd_all = db.query(ICD).filter(ICD.is_deleted == False).all()
    icd_map = {icd.id: icd for icd in icd_all}

    # Group exams by ICD code
    grouped = {}
    for exam, patient, doctor, appt in all_exams:
        if not exam.diagnosis or not isinstance(exam.diagnosis, list):
            continue
        for icd_id in exam.diagnosis:
            if icd_id not in icd_map:
                continue
            icd = icd_map[icd_id]
            code = icd.icd_code
            name = icd.disease_name
            if code not in grouped:
                grouped[code] = {'name': name, 'cases': []}
            grouped[code]['cases'].append({
                'patient': patient.full_name,
                'doctor': doctor.full_name,
                'date': appt.appointment_date.strftime('%d/%m/%Y') if appt.appointment_date else '',
                'time': exam.examination_date.strftime('%H:%M') if exam.examination_date else ''
            })

    # Sort by count desc (match sheet 1 order)
    sorted_codes = sorted(grouped.keys(), key=lambda c: len(grouped[c]['cases']), reverse=True)
    return grouped, sorted_codes


def _write_icd_summary_sheet(from_str, items, mk_border, to_str, total, ws):
    from openpyxl.styles import Font, Alignment, PatternFill
    from datetime import datetime as dt
    period = f"{dt.strptime(from_str,'%Y-%m-%d').strftime('%d/%m/%Y')} – {dt.strptime(to_str,'%Y-%m-%d').strftime('%d/%m/%Y')}"
    ws.append([f"Thống kê bệnh theo ICD — {period}"])
    ws['A1'].font = Font(bold=True, size=13, color="1D4ED8")
    ws.append([])
    ws.append(["STT", "Mã ICD", "Tên bệnh", "Số ca", "Tỷ lệ %"])

    hdr_fill = PatternFill("solid", fgColor="1D4ED8")
    for c in range(1, 6):
        cell = ws.cell(row=3, column=c)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = mk_border()
    ws.row_dimensions[3].height = 22

    z_fill = PatternFill("solid", fgColor="F8FAFC")
    w_fill = PatternFill("solid", fgColor="FFFFFF")
    for i, item in enumerate(items, 1):
        row_i = i + 3
        pct = round(item['count'] / total * 100, 1)
        row_data = [i, item['code'], item['name'], item['count'], pct]
        fill = z_fill if i % 2 == 0 else w_fill
        for c, v in enumerate(row_data, 1):
            cell = ws.cell(row=row_i, column=c, value=v)
            cell.border = mk_border()
            cell.fill = fill
            cell.alignment = Alignment(vertical='center')
        ws.cell(row=row_i, column=4).number_format = '#,##0'
        ws.cell(row=row_i, column=5).number_format = '0.0'

    tr = len(items) + 4
    ws.cell(row=tr, column=3, value="TỔNG CỘNG").font = Font(bold=True)
    ws.cell(row=tr, column=4, value=sum(i['count'] for i in items)).font = Font(bold=True)
    ws.cell(row=tr, column=4).number_format = '#,##0'

    for col, w in [('A',6),('B',12),('C',40),('D',10),('E',10)]:
        ws.column_dimensions[col].width = w
    return period, w_fill, z_fill


@dashboard_bp.route('/api/dashboard/export-icd-excel', methods=['GET'])
@require_auth
@api_error_boundary(error='{error}')
def export_icd_excel(user):
    """Xuất Excel thống kê ICD theo khoảng ngày"""
    import io
    from datetime import date, timedelta, datetime as dt
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Border, Side
        from flask import send_file
    except ImportError:
        return jsonify({'error': 'openpyxl chưa được cài đặt'}), 500

    db = next(get_db())
    try:
        today = date.today()
        from_str = request.args.get('from_date', (today - timedelta(days=6)).isoformat())
        to_str = request.args.get('to_date', today.isoformat())
        start = dt.strptime(from_str, '%Y-%m-%d')
        end = dt.combine(dt.strptime(to_str, '%Y-%m-%d').date(), dt.max.time())

        # Thống kê ICD bằng SQL trực tiếp
        query_sql = """
            SELECT 
                icd.icd_code, 
                icd.disease_name, 
                COUNT(e.id) AS cnt
            FROM examinations e
            JOIN appointments a ON e.appointment_id = a.id
            CROSS JOIN LATERAL jsonb_array_elements_text(e.diagnosis) AS elem
            JOIN icd ON icd.id = elem::int
            WHERE e.diagnosis IS NOT NULL 
              AND e.diagnosis != '[]'::jsonb
              AND a.appointment_date >= :start 
              AND a.appointment_date <= :end
              AND (icd.is_deleted IS NULL OR icd.is_deleted = false)
            GROUP BY icd.icd_code, icd.disease_name
            ORDER BY cnt DESC
        """
        results = db.execute(text(query_sql), {"start": start, "end": end}).fetchall()

        total = sum(r.cnt for r in results) or 1
        items = []
        for r in results:
            items.append({
                'code': r.icd_code,
                'name': r.disease_name,
                'count': r.cnt
            })

        def mk_border():
            s = Side(style='thin', color='D1D5DB')
            return Border(left=s, right=s, top=s, bottom=s)

        wb = openpyxl.Workbook()

        # ===== Sheet 1: Tổng hợp =====
        ws = wb.active
        ws.title = "Tổng hợp"

        period, w_fill, z_fill = _write_icd_summary_sheet(from_str, items, mk_border, to_str, total, ws)

        # ===== Sheet 2: Chi tiết ca khám nhóm theo ICD =====

        grouped, sorted_codes = _group_icd_cases(db, end, start)

        ws2 = wb.create_sheet("Chi tiết ca khám")
        ws2.append([f"Chi tiết ca khám theo ICD — {period}"])
        ws2['A1'].font = Font(bold=True, size=13, color="7C3AED")
        ws2.append([])

        row_num = 3
        grp_fill = PatternFill("solid", fgColor="EDE9FE")
        detail_hdr_fill = PatternFill("solid", fgColor="F5F3FF")

        _write_icd_case_groups(detail_hdr_fill, grouped, grp_fill, mk_border, row_num, sorted_codes, w_fill, ws2, z_fill)

        for col, w in [('A',6),('B',25),('C',30),('D',12),('E',8)]:
            ws2.column_dimensions[col].width = w

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        fname = f"icd_{from_str}_{to_str}.xlsx"
        return send_file(output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name=fname)

    finally:
        db.close()

@dashboard_bp.route('/api/dashboard/icd-detail', methods=['GET'])
@require_auth
@api_error_boundary(error='{error}', items=[])
def get_icd_detail(user):
    """Chi tiết các ca khám theo mã ICD trong khoảng ngày"""
    db = next(get_db())
    try:
        from datetime import date, timedelta, datetime as dt
        from app.models.patient import Patient
        from app.models.user import User

        icd_code = request.args.get('icd_code', '')
        today = date.today()
        from_str = request.args.get('from_date', (today - timedelta(days=6)).isoformat())
        to_str = request.args.get('to_date', today.isoformat())
        start = dt.strptime(from_str, '%Y-%m-%d')
        end = dt.combine(dt.strptime(to_str, '%Y-%m-%d').date(), dt.max.time())

        from app.models.icd import ICD
        # 1. Tìm ICD theo icd_code
        icd_rec = db.query(ICD).filter(ICD.icd_code == icd_code, ICD.is_deleted == False).first() if icd_code != '—' else None

        # 2. Query ca khám
        query = db.query(Examination, Patient, User, Appointment)\
                  .join(Appointment, Examination.appointment_id == Appointment.id)\
                  .join(Patient, Examination.patient_id == Patient.id)\
                  .join(User, Examination.doctor_id == User.id)\
                  .filter(
                      Examination.diagnosis.isnot(None),
                      Appointment.appointment_date >= start,
                      Appointment.appointment_date <= end
                  )
                  
        if icd_code != '—' and icd_rec:
            # Lọc các bản ghi chứa icd_rec.id trong mảng JSONB diagnosis
            query = query.filter(Examination.diagnosis.contains([icd_rec.id]))
            
        rows = query.order_by(Appointment.appointment_date.desc()).all()

        items = []
        from app.utils.examination_utils import resolve_diagnosis_to_str
        for exam, patient, doctor, appt in rows:
            diag_str = resolve_diagnosis_to_str(db, exam.diagnosis)
            items.append({
                'patient_name': patient.full_name,
                'phone': patient.phone_number if hasattr(patient, 'phone_number') else '',
                'doctor_name': doctor.full_name,
                'exam_date': appt.appointment_date.strftime('%d/%m/%Y') if appt.appointment_date else '',
                'diagnosis': diag_str,
                'exam_time': exam.examination_date.strftime('%H:%M') if exam.examination_date else ''
            })

        return jsonify({'items': items, 'icd_code': icd_code}), 200

    finally:
        db.close()

# Route/hàm còn lại nằm ở dashboard_revenue.py; import để đăng ký route và giữ tên cũ trên module này.
from app.api.dashboard_revenue import (  # noqa: E402,F401
    export_dashboard_excel,
    get_revenue_detail,
    get_revenue,
    get_exam_today,
    get_top_icd,
)

# Route/hàm còn lại nằm ở dashboard_referral_staff.py; import để đăng ký route và giữ tên cũ trên module này.
from app.api.dashboard_referral_staff import (  # noqa: E402,F401
    get_referral_sources,
    get_referral_source_detail,
    export_referral_source_excel,
    get_staff_online,
)

# Route/hàm còn lại nằm ở dashboard_income_exam_stats.py; import để đăng ký route và giữ tên cũ trên module này.
from app.api.dashboard_income_exam_stats import (  # noqa: E402,F401
    export_thu_chi_excel,
    get_exam_stats_by_day,
    get_exam_detail_by_day,
)
