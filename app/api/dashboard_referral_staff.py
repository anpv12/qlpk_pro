"""app.api.dashboard: phần 3 — tách từ dashboard.py (import ở cuối dashboard.py để đăng ký route/giữ tên cũ)."""

from flask import jsonify, request
from app.core.database import get_db
from app.models.user import User, UserRole
from app.api.auth import require_auth
from app.realtime.presence import get_online_user_ids
from app.utils.referral_source import REFERRAL_SOURCE_BY_KEY, REFERRAL_SOURCE_CONFIG
from app.utils.upload_storage import normalize_upload_url
from app.api.dashboard import (  # noqa: E402 — module gốc đã khởi tạo xong các tên này
    _get_referral_source_cases,
    _parse_referral_source_date_range,
    dashboard_bp,
    logger,
)


@dashboard_bp.route('/api/dashboard/referral-sources', methods=['GET'])
@require_auth
def get_referral_sources(user):
    """Thống kê lượt khám theo nguồn giới thiệu của bệnh nhân."""
    db = next(get_db())
    try:
        from_str, to_str, start, end = _parse_referral_source_date_range()
        cases = _get_referral_source_cases(db, start, end)

        counts = {item['key']: 0 for item in REFERRAL_SOURCE_CONFIG}
        for case in cases:
            counts[case['source_key']] += 1

        items = []
        for item in REFERRAL_SOURCE_CONFIG:
            items.append({
                'key': item['key'],
                'label': item['label'],
                'color': item['color'],
                'count': counts[item['key']],
            })

        return jsonify({
            'from_date': from_str,
            'to_date': to_str,
            'total': sum(counts.values()),
            'items': items,
        }), 200

    except Exception as e:
        logger.error(f"Error getting referral source stats: {e}", exc_info=True)
        return jsonify({'error': str(e), 'items': []}), 500
    finally:
        db.close()


@dashboard_bp.route('/api/dashboard/referral-source-detail', methods=['GET'])
@require_auth
def get_referral_source_detail(user):
    """Danh sách ca khám thuộc một nguồn giới thiệu."""
    db = next(get_db())
    try:
        source_key = request.args.get('source', '').strip()
        if source_key not in REFERRAL_SOURCE_BY_KEY:
            return jsonify({'error': 'Nguồn giới thiệu không hợp lệ', 'items': []}), 400

        from_str, to_str, start, end = _parse_referral_source_date_range()
        cases = [
            case for case in _get_referral_source_cases(db, start, end)
            if case['source_key'] == source_key
        ]

        source_cfg = REFERRAL_SOURCE_BY_KEY[source_key]
        return jsonify({
            'from_date': from_str,
            'to_date': to_str,
            'source': source_key,
            'source_label': source_cfg['label'],
            'items': cases,
            'total': len(cases),
        }), 200

    except Exception as e:
        logger.error(f"Error getting referral source detail: {e}", exc_info=True)
        return jsonify({'error': str(e), 'items': []}), 500
    finally:
        db.close()


def _write_referral_detail_sheet(detail_hdr_fill, group_fill, grouped_cases, mk_border, period, w_fill, wb, z_fill):
    from openpyxl.styles import Font, Alignment
    ws2 = wb.create_sheet('Chi tiết')
    ws2.append([f'Chi tiết lượt khám theo nguồn giới thiệu — {period}'])
    ws2['A1'].font = Font(bold=True, size=13, color='0F766E')
    ws2.append([])

    row_num = 3
    detail_headers = ['STT', 'Bệnh nhân', 'Điện thoại', 'Ngày khám', 'Giờ', 'Lý do khám', 'Nguồn nhập']
    for cfg in REFERRAL_SOURCE_CONFIG:
        source_cases = grouped_cases[cfg['key']]
        group_title = f"{cfg['label']} — {len(source_cases)} lượt khám"
        for col in range(1, 8):
            cell = ws2.cell(row=row_num, column=col, value=group_title if col == 1 else '')
            cell.font = Font(bold=True, size=11, color='0F766E')
            cell.fill = group_fill
            cell.border = mk_border()
            cell.alignment = Alignment(vertical='center')
        ws2.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=7)
        row_num += 1

        for col, value in enumerate(detail_headers, 1):
            cell = ws2.cell(row=row_num, column=col, value=value)
            cell.font = Font(bold=True, size=10, color='64748B')
            cell.fill = detail_hdr_fill
            cell.border = mk_border()
            cell.alignment = Alignment(horizontal='center', vertical='center')
        row_num += 1

        if source_cases:
            for index, case in enumerate(source_cases, 1):
                fill = z_fill if index % 2 == 0 else w_fill
                row_values = [
                    index,
                    case['patient_name'],
                    case['phone'],
                    case['exam_date'],
                    case['exam_time'],
                    case['main_reason'] or '—',
                    case['source_value'] or case['source_label'],
                ]
                for col, value in enumerate(row_values, 1):
                    cell = ws2.cell(row=row_num, column=col, value=value)
                    cell.border = mk_border()
                    cell.fill = fill
                    cell.alignment = Alignment(vertical='top', wrap_text=col in (6, 7))
                row_num += 1
        else:
            for col in range(1, 8):
                cell = ws2.cell(row=row_num, column=col, value='Không có dữ liệu' if col == 2 else '')
                cell.border = mk_border()
                cell.fill = w_fill
                cell.alignment = Alignment(vertical='center')
            row_num += 1

        row_num += 1

    for col, width in [('A', 8), ('B', 26), ('C', 16), ('D', 14), ('E', 10), ('F', 44), ('G', 30)]:
        ws2.column_dimensions[col].width = width


def _write_referral_summary_sheet(counts, hdr_fill, mk_border, pct_base, period, total, w_fill, wb, z_fill):
    from openpyxl.styles import Font, Alignment
    ws = wb.active
    ws.title = 'Tổng hợp'
    ws.append([f'Thống kê nguồn giới thiệu — {period}'])
    ws['A1'].font = Font(bold=True, size=13, color='0F766E')
    ws.append([])
    ws.append(['STT', 'Nguồn giới thiệu', 'Số lượt khám', 'Tỷ lệ %'])

    for col in range(1, 5):
        cell = ws.cell(row=3, column=col)
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = hdr_fill
        cell.border = mk_border()
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for index, cfg in enumerate(REFERRAL_SOURCE_CONFIG, 1):
        row_idx = index + 3
        count = counts[cfg['key']]
        percentage = round(count / pct_base * 100, 1)
        fill = z_fill if index % 2 == 0 else w_fill
        for col, value in enumerate([index, cfg['label'], count, percentage], 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = mk_border()
            cell.fill = fill
            cell.alignment = Alignment(vertical='center')
        ws.cell(row=row_idx, column=3).number_format = '#,##0'
        ws.cell(row=row_idx, column=4).number_format = '0.0'

    total_row = len(REFERRAL_SOURCE_CONFIG) + 4
    ws.cell(row=total_row, column=2, value='TỔNG CỘNG').font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=total).font = Font(bold=True)
    ws.cell(row=total_row, column=3).number_format = '#,##0'
    for col in range(1, 5):
        ws.cell(row=total_row, column=col).border = mk_border()

    for col, width in [('A', 8), ('B', 34), ('C', 14), ('D', 12)]:
        ws.column_dimensions[col].width = width


@dashboard_bp.route('/api/dashboard/export-referral-source-excel', methods=['GET'])
@require_auth
def export_referral_source_excel(user):
    """Xuất Excel thống kê nguồn giới thiệu theo khoảng ngày."""
    import io
    from datetime import datetime as dt
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Border, Side
        from flask import send_file
    except ImportError:
        return jsonify({'error': 'openpyxl chưa được cài đặt'}), 500

    db = next(get_db())
    try:
        from_str, to_str, start, end = _parse_referral_source_date_range()
        cases = _get_referral_source_cases(db, start, end)

        counts = {item['key']: 0 for item in REFERRAL_SOURCE_CONFIG}
        grouped_cases = {item['key']: [] for item in REFERRAL_SOURCE_CONFIG}
        for case in cases:
            counts[case['source_key']] += 1
            grouped_cases[case['source_key']].append(case)

        total = sum(counts.values())
        pct_base = total or 1
        period = f"{dt.strptime(from_str, '%Y-%m-%d').strftime('%d/%m/%Y')} – {dt.strptime(to_str, '%Y-%m-%d').strftime('%d/%m/%Y')}"

        def mk_border():
            side = Side(style='thin', color='D1D5DB')
            return Border(left=side, right=side, top=side, bottom=side)

        wb = openpyxl.Workbook()
        hdr_fill = PatternFill('solid', fgColor='0F766E')
        group_fill = PatternFill('solid', fgColor='CCFBF1')
        detail_hdr_fill = PatternFill('solid', fgColor='F0FDF4')
        z_fill = PatternFill('solid', fgColor='F8FAFC')
        w_fill = PatternFill('solid', fgColor='FFFFFF')

        _write_referral_summary_sheet(counts, hdr_fill, mk_border, pct_base, period, total, w_fill, wb, z_fill)

        _write_referral_detail_sheet(detail_hdr_fill, group_fill, grouped_cases, mk_border, period, w_fill, wb, z_fill)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"nguon_gioi_thieu_{from_str}_{to_str}.xlsx"
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )

    except Exception as e:
        logger.error(f"Export referral source excel error: {e}", exc_info=True)
        return jsonify({'error': str(e)}), 500
    finally:
        db.close()


@dashboard_bp.route('/api/dashboard/staff-online', methods=['GET'])
@require_auth
def get_staff_online(user):
    """Trạng thái online của bác sĩ / tâm lý gia / lễ tân"""
    db = next(get_db())
    try:
        online_user_ids = get_online_user_ids()

        staff = db.query(User).filter(
            User.role.in_([UserRole.DOCTOR, UserRole.PSYCHOLOGIST, UserRole.STAFF]),
            User.is_active == True
        ).order_by(User.full_name).all()

        role_map = {
            UserRole.DOCTOR: 'Bác sĩ',
            UserRole.PSYCHOLOGIST: 'Tâm lý gia',
            UserRole.STAFF: 'Lễ tân'
        }

        items = []
        for s in staff:
            is_online = s.id in online_user_ids
            role_label = role_map.get(s.role, 'Nhân viên')
            items.append({
                'id': s.id,
                'full_name': s.full_name,
                'role': str(s.role.value) if s.role else '',
                'role_label': role_label,
                'avatar': normalize_upload_url(s.avatar),
                'is_online': is_online,
                'last_login': s.last_login.isoformat() if s.last_login else None
            })

        online_count = sum(1 for i in items if i['is_online'])
        return jsonify({
            'items': items,
            'online_count': online_count,
            'total': len(items)
        }), 200

    except Exception as e:
        logger.error(f"Error getting staff online: {e}")
        return jsonify({'error': str(e), 'items': []}), 500
    finally:
        db.close()
