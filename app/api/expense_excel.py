"""expense helpers split out by topic (excel); re-exported by app.api.expense."""

import logging
from flask import request
from app.models.expense import Expense
from app.models.expense_column import ExpenseColumn
from datetime import datetime

logger = logging.getLogger('app.api.expense')


def _write_expense_rows(center, col_map, data_font, expense_dicts, left_align, right_align, thin_border, ws):
    # Data rows
    for ri, exp in enumerate(expense_dicts):
        row_num = ri + 2
        # STT
        stt_cell = ws.cell(row=row_num, column=1, value=ri + 1)
        stt_cell.font = data_font
        stt_cell.alignment = center
        stt_cell.border = thin_border

        for excel_col, col_cfg in col_map:
            val = exp.get(col_cfg.col_id, '')
            cell = ws.cell(row=row_num, column=excel_col, value=val)
            cell.font = data_font
            cell.border = thin_border

            if col_cfg.col_type == 'number':
                cell.alignment = right_align
                cell.number_format = '#,##0'
                try:
                    cell.value = float(val) if val else 0
                except (ValueError, TypeError) as exc:
                    logger.debug("Giữ nguyên giá trị ô Excel không chuyển đổi được: %s", exc)
            elif col_cfg.col_type == 'date':
                cell.alignment = center
                # Chuyển string dd/mm/yyyy → Python date để Excel nhận thật sự là ngày
                if val and isinstance(val, str):
                    try:
                        cell.value = datetime.strptime(val, '%d/%m/%Y').date()
                    except ValueError as exc:
                        logger.debug("Giữ nguyên giá trị ô Excel không chuyển đổi được: %s", exc)
                cell.number_format = 'DD/MM/YYYY'
            else:
                cell.alignment = left_align


def _load_expense_export_data(db):
    # Lấy cấu hình cột active
    col_configs = db.query(ExpenseColumn).filter(
        ExpenseColumn.is_deleted == False
    ).order_by(ExpenseColumn.position).all()

    # Fallback về cột mặc định nếu DB chưa có cấu hình (giống FE DEFAULT_COLUMNS)
    if not col_configs:
        from types import SimpleNamespace
        DEFAULT_COLS = [
            SimpleNamespace(col_id='date',     name='Ngày',      col_type='date',   width=70),
            SimpleNamespace(col_id='desc',     name='Nội dung',  col_type='text',   width=200),
            SimpleNamespace(col_id='category', name='Hạng mục',  col_type='text',   width=120),
            SimpleNamespace(col_id='type',     name='Loại chi',  col_type='text',   width=120),
            SimpleNamespace(col_id='amount',   name='Số tiền',   col_type='number', width=90),
            SimpleNamespace(col_id='payment',  name='Hình thức', col_type='text',   width=100),
            SimpleNamespace(col_id='note',     name='Nguồn tiền',col_type='text',   width=120),
        ]
        col_configs = DEFAULT_COLS

    # Lấy expenses theo date range
    query = db.query(Expense)
    date_from = request.args.get('from')
    date_to = request.args.get('to')
    if date_from:
        parsed = parse_date(date_from)
        if parsed:
            query = query.filter(Expense.date >= parsed)
    if date_to:
        parsed = parse_date(date_to)
        if parsed:
            query = query.filter(Expense.date <= parsed)
    expenses = query.order_by(Expense.date.desc()).all()
    expense_dicts = [e.to_dict() for e in expenses]
    return col_configs, expense_dicts


def _write_expense_total_row(col_configs, col_map, expense_dicts, left_align, right_align, thin_border, total_fill, total_font, ws):
    from openpyxl.utils import get_column_letter
    # Total row
    total_row = len(expense_dicts) + 2
    ws.cell(row=total_row, column=1, value='').border = thin_border
    ws.cell(row=total_row, column=1).fill = total_fill
    for excel_col, col_cfg in col_map:
        cell = ws.cell(row=total_row, column=excel_col)
        cell.border = thin_border
        cell.fill = total_fill
        cell.font = total_font
        if col_cfg.col_type == 'number':
            col_letter = get_column_letter(excel_col)
            cell.value = f'=SUM({col_letter}2:{col_letter}{total_row - 1})'
            cell.number_format = '#,##0'
            cell.alignment = right_align
        elif col_cfg.col_id == 'desc':
            cell.value = f'{len(expense_dicts)} khoản'
            cell.alignment = left_align

    # Freeze & filter
    ws.freeze_panes = 'A2'
    last_col = get_column_letter(len(col_configs) + 1)
    ws.auto_filter.ref = f'A1:{last_col}{total_row}'
    ws.sheet_properties.tabColor = '0F766E'


def parse_date(date_str):
    """Parse date from dd/mm/yyyy or dd/mm format"""
    if not date_str:
        return None
    for fmt in ('%d/%m/%Y', '%d/%m'):
        try:
            d = datetime.strptime(date_str.strip(), fmt)
            if fmt == '%d/%m':
                d = d.replace(year=datetime.now().year)
            return d.date()
        except ValueError:
            continue
    return None
