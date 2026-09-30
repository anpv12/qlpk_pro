from flask import Blueprint, request, jsonify, send_file
from app.models.expense import Expense
from app.models.expense_column import ExpenseColumn
from app.core.database import get_db
from app.api.auth import require_auth
from app.realtime.events import emit_finance_changed
from datetime import datetime
import logging
import io
from app.utils.api_error_contract import api_error_boundary

logger = logging.getLogger(__name__)

expense_bp = Blueprint('expense', __name__, url_prefix='/api/expenses')


@expense_bp.route('/template', methods=['GET'])
@require_auth
def download_template(user):
    """Tạo file Excel mẫu import chi tiêu"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = 'Mẫu chi tiêu'

    # Column config
    cols = [
        {'name': 'Ngày', 'width': 14, 'sample': ['14/03/2026', '14/03/2026']},
        {'name': 'Nội dung', 'width': 30, 'sample': ['Mua văn phòng phẩm', 'Tiền điện tháng 3']},
        {'name': 'Hạng mục', 'width': 20, 'sample': ['Chi phát sinh', 'Chi thường xuyên'],
         'options': ['Chi phát sinh', 'Chi thường xuyên', 'Chi vật tư tiêu hao']},
        {'name': 'Loại chi', 'width': 20, 'sample': ['Mua sắm vật tư', 'Tiền điện'],
         'options': ['Thuê nhà','Tiền điện','Nước','Rác','Tiền lương','Sửa chữa vật tư','Mua sắm vật tư','Khác','Từ thiện','Quan hệ','Hàng ngày','Hợp đồng']},
        {'name': 'Số tiền', 'width': 15, 'sample': [350000, 2800000]},
        {'name': 'Hình thức', 'width': 16, 'sample': ['Tiền mặt', 'Chuyển khoản'],
         'options': ['Tiền mặt', 'Chuyển khoản']},
        {'name': 'Nguồn tiền', 'width': 16, 'sample': ['Tiền lễ tân', 'Tài khoản PK'],
         'options': ['BS Hiến','BS Tuấn','TLG Khương','Tươi','Tiền lễ tân','Tài khoản PK','Tiền khám']},
    ]

    # Styles
    header_fill = PatternFill(start_color='1B4F72', end_color='1B4F72', fill_type='solid')
    header_font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    data_font = Font(name='Arial', size=11)
    thin_border = Border(
        left=Side(style='thin', color='D5D8DC'),
        right=Side(style='thin', color='D5D8DC'),
        top=Side(style='thin', color='D5D8DC'),
        bottom=Side(style='thin', color='D5D8DC'),
    )
    center = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    # Header row
    for ci, col in enumerate(cols, 1):
        cell = ws.cell(row=1, column=ci, value=col['name'])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(ci)].width = col['width']

    # Sample data rows
    alt_fill = PatternFill(start_color='EBF5FB', end_color='EBF5FB', fill_type='solid')
    for ri in range(2):
        for ci, col in enumerate(cols, 1):
            cell = ws.cell(row=ri + 2, column=ci, value=col['sample'][ri])
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center if ci in (1, 5) else left_align
            if ri % 2 == 1:
                cell.fill = alt_fill
            if col['name'] == 'Số tiền':
                cell.number_format = '#,##0'

    # Data validation dropdowns (apply to rows 2-1000)
    for ci, col in enumerate(cols, 1):
        if 'options' in col:
            formula = '"' + ','.join(col['options']) + '"'
            dv = DataValidation(type='list', formula1=formula, allow_blank=True)
            dv.error = f'Chọn giá trị từ danh sách {col["name"]}'
            dv.errorTitle = 'Giá trị không hợp lệ'
            dv.prompt = f'Chọn {col["name"]}'
            dv.promptTitle = col['name']
            col_letter = get_column_letter(ci)
            dv.add(f'{col_letter}2:{col_letter}1000')
            ws.add_data_validation(dv)

    # Freeze header row
    ws.freeze_panes = 'A2'
    ws.sheet_properties.tabColor = '1B4F72'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='mau-import-chi-tieu.xlsx'
    )


def _write_expense_rows(center, col_map, data_font, expense_dicts, left_align, right_align, thin_border, ws):
    # Data rows
    for ri, exp in enumerate(expense_dicts):
        row_num = ri + 2
        # STT
        stt_cell = ws.cell(row=row_num, column=1, value=ri + 1)
        stt_cell.font = data_font
        stt_cell.alignment = center
        stt_cell.border = thin_border

        for excel_col, col_cfg in col_map:
            val = exp.get(col_cfg.col_id, '')
            cell = ws.cell(row=row_num, column=excel_col, value=val)
            cell.font = data_font
            cell.border = thin_border

            if col_cfg.col_type == 'number':
                cell.alignment = right_align
                cell.number_format = '#,##0'
                try:
                    cell.value = float(val) if val else 0
                except (ValueError, TypeError) as exc:
                    logger.debug("Giữ nguyên giá trị ô Excel không chuyển đổi được: %s", exc)
            elif col_cfg.col_type == 'date':
                cell.alignment = center
                # Chuyển string dd/mm/yyyy → Python date để Excel nhận thật sự là ngày
                if val and isinstance(val, str):
                    try:
                        cell.value = datetime.strptime(val, '%d/%m/%Y').date()
                    except ValueError as exc:
                        logger.debug("Giữ nguyên giá trị ô Excel không chuyển đổi được: %s", exc)
                cell.number_format = 'DD/MM/YYYY'
            else:
                cell.alignment = left_align


def _load_expense_export_data(db):
    # Lấy cấu hình cột active
    col_configs = db.query(ExpenseColumn).filter(
        ExpenseColumn.is_deleted == False
    ).order_by(ExpenseColumn.position).all()

    # Fallback về cột mặc định nếu DB chưa có cấu hình (giống FE DEFAULT_COLUMNS)
    if not col_configs:
        from types import SimpleNamespace
        DEFAULT_COLS = [
            SimpleNamespace(col_id='date',     name='Ngày',      col_type='date',   width=70),
            SimpleNamespace(col_id='desc',     name='Nội dung',  col_type='text',   width=200),
            SimpleNamespace(col_id='category', name='Hạng mục',  col_type='text',   width=120),
            SimpleNamespace(col_id='type',     name='Loại chi',  col_type='text',   width=120),
            SimpleNamespace(col_id='amount',   name='Số tiền',   col_type='number', width=90),
            SimpleNamespace(col_id='payment',  name='Hình thức', col_type='text',   width=100),
            SimpleNamespace(col_id='note',     name='Nguồn tiền',col_type='text',   width=120),
        ]
        col_configs = DEFAULT_COLS

    # Lấy expenses theo date range
    query = db.query(Expense)
    date_from = request.args.get('from')
    date_to = request.args.get('to')
    if date_from:
        parsed = parse_date(date_from)
        if parsed:
            query = query.filter(Expense.date >= parsed)
    if date_to:
        parsed = parse_date(date_to)
        if parsed:
            query = query.filter(Expense.date <= parsed)
    expenses = query.order_by(Expense.date.desc()).all()
    expense_dicts = [e.to_dict() for e in expenses]
    return col_configs, expense_dicts


def _write_expense_total_row(col_configs, col_map, expense_dicts, left_align, right_align, thin_border, total_fill, total_font, ws):
    from openpyxl.utils import get_column_letter
    # Total row
    total_row = len(expense_dicts) + 2
    ws.cell(row=total_row, column=1, value='').border = thin_border
    ws.cell(row=total_row, column=1).fill = total_fill
    for excel_col, col_cfg in col_map:
        cell = ws.cell(row=total_row, column=excel_col)
        cell.border = thin_border
        cell.fill = total_fill
        cell.font = total_font
        if col_cfg.col_type == 'number':
            col_letter = get_column_letter(excel_col)
            cell.value = f'=SUM({col_letter}2:{col_letter}{total_row - 1})'
            cell.number_format = '#,##0'
            cell.alignment = right_align
        elif col_cfg.col_id == 'desc':
            cell.value = f'{len(expense_dicts)} khoản'
            cell.alignment = left_align

    # Freeze & filter
    ws.freeze_panes = 'A2'
    last_col = get_column_letter(len(col_configs) + 1)
    ws.auto_filter.ref = f'A1:{last_col}{total_row}'
    ws.sheet_properties.tabColor = '0F766E'


@expense_bp.route('/export', methods=['GET'])
@require_auth
@api_error_boundary(detail='Lỗi xuất Excel')
def export_excel(user):
    """Xuất Excel chi tiêu có format"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    db = next(get_db())
    try:
        col_configs, expense_dicts = _load_expense_export_data(db)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Chi tiêu'

        # Styles
        header_fill = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
        header_font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        data_font = Font(name='Arial', size=10)
        total_font = Font(name='Arial', bold=True, size=11, color='0F766E')
        thin_border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000'),
        )
        center = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
        right_align = Alignment(horizontal='right', vertical='center')
        total_fill = PatternFill(start_color='F0FDF4', end_color='F0FDF4', fill_type='solid')

        # STT column + dynamic columns
        ws.cell(row=1, column=1, value='STT').font = header_font
        ws.cell(row=1, column=1).fill = header_fill
        ws.cell(row=1, column=1).alignment = center
        ws.cell(row=1, column=1).border = thin_border
        ws.column_dimensions['A'].width = 6

        col_map = []  # (col_index_in_excel, col_config)
        for ci, col_cfg in enumerate(col_configs):
            excel_col = ci + 2  # offset by STT column
            cell = ws.cell(row=1, column=excel_col, value=col_cfg.name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin_border

            width = col_cfg.width or 100
            ws.column_dimensions[get_column_letter(excel_col)].width = max(width / 7, 12)
            col_map.append((excel_col, col_cfg))

        _write_expense_rows(center, col_map, data_font, expense_dicts, left_align, right_align, thin_border, ws)

        _write_expense_total_row(col_configs, col_map, expense_dicts, left_align, right_align, thin_border, total_fill, total_font, ws)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'chi-tieu-{datetime.now().strftime("%d-%m-%Y")}.xlsx'
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename
        )
    finally:
        db.close()

def parse_date(date_str):
    """Parse date from dd/mm/yyyy or dd/mm format"""
    if not date_str:
        return None
    for fmt in ('%d/%m/%Y', '%d/%m'):
        try:
            d = datetime.strptime(date_str.strip(), fmt)
            if fmt == '%d/%m':
                d = d.replace(year=datetime.now().year)
            return d.date()
        except ValueError:
            continue
    return None


@expense_bp.route('', methods=['GET'])
@require_auth
@api_error_boundary(detail='Lỗi tải danh sách chi tiêu')
def list_expenses(user):
    """Lấy danh sách khoản chi, hỗ trợ filter theo ngày"""
    db = next(get_db())
    try:
        query = db.query(Expense)

        date_from = request.args.get('date_from')
        date_to = request.args.get('date_to')

        if date_from:
            d = parse_date(date_from)
            if d:
                query = query.filter(Expense.date >= d)
        if date_to:
            d = parse_date(date_to)
            if d:
                query = query.filter(Expense.date <= d)

        expenses = query.order_by(Expense.date.desc(), Expense.id.desc()).all()
        return jsonify([e.to_dict() for e in expenses]), 200
    finally:
        db.close()


@expense_bp.route('', methods=['POST'])
@require_auth
@api_error_boundary(detail='Lỗi tạo khoản chi')
def create_expense(user):
    """Tạo khoản chi mới"""
    db = next(get_db())
    try:
        data = request.get_json()
        expense = Expense(
            date=parse_date(data.get('date')) or datetime.now().date(),
            type=data.get('type', ''),
            category=data.get('category', ''),
            description=data.get('desc', ''),
            person=data.get('person', ''),
            amount=float(data.get('amount', 0)) if data.get('amount') else 0,
            payment=data.get('payment', ''),
            note=data.get('note', ''),
            created_by=user.id
        )
        db.add(expense)
        db.commit()
        db.refresh(expense)
        emit_finance_changed('expense_created', entity='expense', entity_id=expense.id)
        return jsonify(expense.to_dict()), 201
    finally:
        db.close()


@expense_bp.route('/<int:expense_id>', methods=['PUT'])
@require_auth
@api_error_boundary(detail='Lỗi cập nhật khoản chi')
def update_expense(user, expense_id):
    """Cập nhật khoản chi"""
    db = next(get_db())
    try:
        expense = db.query(Expense).filter(Expense.id == expense_id).first()
        if not expense:
            return jsonify({'detail': 'Không tìm thấy khoản chi'}), 404

        data = request.get_json()

        # Known DB fields mapping
        KNOWN_FIELDS = {
            'date': lambda v: setattr(expense, 'date', parse_date(v) or expense.date),
            'type': lambda v: setattr(expense, 'type', v),
            'category': lambda v: setattr(expense, 'category', v),
            'desc': lambda v: setattr(expense, 'description', v),
            'person': lambda v: setattr(expense, 'person', v),
            'amount': lambda v: setattr(expense, 'amount', float(v) if v else 0),
            'payment': lambda v: setattr(expense, 'payment', v),
            'note': lambda v: setattr(expense, 'note', v),
        }

        extra = dict(expense.extra_data or {})
        for key, val in data.items():
            if key in ('id', 'created_by', 'created_at', 'updated_at', 'created_by_name'):
                continue
            if key in KNOWN_FIELDS:
                KNOWN_FIELDS[key](val)
            else:
                extra[key] = val
        expense.extra_data = extra

        expense.updated_at = datetime.now()
        db.commit()
        db.refresh(expense)
        emit_finance_changed('expense_updated', entity='expense', entity_id=expense.id)
        return jsonify(expense.to_dict()), 200
    finally:
        db.close()


@expense_bp.route('/<int:expense_id>', methods=['DELETE'])
@require_auth
@api_error_boundary(detail='Lỗi xóa khoản chi')
def delete_expense(user, expense_id):
    """Xóa khoản chi"""
    db = next(get_db())
    try:
        expense = db.query(Expense).filter(Expense.id == expense_id).first()
        if not expense:
            return jsonify({'detail': 'Không tìm thấy khoản chi'}), 404

        db.delete(expense)
        db.commit()
        emit_finance_changed('expense_deleted', entity='expense', entity_id=expense_id)
        return jsonify({'message': 'Đã xóa khoản chi'}), 200
    finally:
        db.close()


@expense_bp.route('/bulk', methods=['POST'])
@require_auth
@api_error_boundary(detail='Lỗi import chi tiêu')
def bulk_create_expenses(user):
    """Import nhiều khoản chi cùng lúc (cho CSV import)"""
    db = next(get_db())
    try:
        data = request.get_json()
        items = data.get('items', [])
        if not items:
            return jsonify({'detail': 'Không có dữ liệu import'}), 400

        created = []
        for item in items:
            expense = Expense(
                date=parse_date(item.get('date')) or datetime.now().date(),
                type=item.get('type', ''),
                category=item.get('category', ''),
                description=item.get('desc', ''),
                person=item.get('person', ''),
                amount=float(item.get('amount', 0)) if item.get('amount') else 0,
                payment=item.get('payment', ''),
                note=item.get('note', ''),
                created_by=user.id
            )
            db.add(expense)
            created.append(expense)

        db.commit()
        for e in created:
            db.refresh(e)
        emit_finance_changed('expenses_imported', entity='expense', extra={'count': len(created)})

        return jsonify({
            'message': f'Đã import {len(created)} khoản chi',
            'items': [e.to_dict() for e in created]
        }), 201
    finally:
        db.close()


@expense_bp.route('/columns', methods=['GET'])
@require_auth
def get_columns(user):
    """Lấy cấu hình cột active (shared, không phân biệt user)"""
    db = next(get_db())
    try:
        cols = db.query(ExpenseColumn).filter(
            ExpenseColumn.is_deleted == False
        ).order_by(ExpenseColumn.position).all()
        return jsonify([c.to_dict() for c in cols])
    finally:
        db.close()


@expense_bp.route('/columns', methods=['POST'])
@require_auth
@api_error_boundary(detail='Lỗi lưu cấu hình cột')
def save_columns(user):
    """Lưu cấu hình cột (upsert + soft-delete)"""
    db = next(get_db())
    try:
        data = request.get_json()
        if not isinstance(data, list):
            return jsonify({'detail': 'Dữ liệu phải là mảng'}), 400

        incoming_ids = set()
        for i, col in enumerate(data):
            col_id = col.get('id', f'col_{i}')
            incoming_ids.add(col_id)

            existing = db.query(ExpenseColumn).filter(
                ExpenseColumn.col_id == col_id
            ).first()

            if existing:
                existing.name = col.get('name', existing.name)
                existing.col_type = col.get('type', existing.col_type)
                existing.options = col.get('options') if col.get('options') else None
                existing.formula = col.get('formula') or None
                existing.width = col.get('width', 100)
                existing.position = i
                existing.is_deleted = False
                existing.deleted_at = None
            else:
                new_col = ExpenseColumn(
                    col_id=col_id,
                    name=col.get('name', ''),
                    col_type=col.get('type', 'text'),
                    options=col.get('options') if col.get('options') else None,
                    formula=col.get('formula') or None,
                    width=col.get('width', 100),
                    position=i,
                    created_by=user.id,
                )
                db.add(new_col)

        # Soft-delete cột không còn trong danh sách
        removed = db.query(ExpenseColumn).filter(
            ExpenseColumn.col_id.notin_(incoming_ids),
            ExpenseColumn.is_deleted == False
        ).all()
        for col in removed:
            col.is_deleted = True
            col.deleted_at = datetime.now()

        db.commit()
        emit_finance_changed('expense_columns_saved', entity='expense_column', extra={
            'count': len(data),
            'removed_count': len(removed),
        })
        return jsonify({'message': f'Đã lưu {len(data)} cột'})
    finally:
        db.close()
